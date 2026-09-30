"""Экспорт нагрузки планирования в CSV / XLSX."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.planning_schemas import PlanningWorkloadOut

MOSCOW_TZ = ZoneInfo("Europe/Moscow")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="CCCCCC")
_HEADER_FONT = Font(bold=True)
_TEXT_ALIGNMENT = Alignment(wrap_text=True, vertical="top")

EXPORT_HEADERS = [
    "ФИО",
    "Отделы",
    "Экспертиза",
    "Строка",
    "Номер проекта",
    "Название проекта",
    "День",
    "План ч",
    "Факт ч",
    "Ёмкость ч",
    "Свободно ч",
    "Отсутствие",
    "Рабочий день",
]


def _dec(value: Decimal | int | float | None) -> str:
    if value is None:
        return "0"
    text = format(Decimal(str(value)), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _join(parts: list[str]) -> str:
    return ", ".join(part for part in parts if part)


def _iter_export_rows(workload: PlanningWorkloadOut, *, view_mode: str) -> list[list[str]]:
    rows: list[list[str]] = []
    by_project = view_mode == "byProject"
    for employee in workload.employees:
        departments = _join(employee.departmentNames)
        expertises = _join(employee.expertises)
        for day in workload.days:
            day_key = day.isoformat()
            cell = employee.days.get(day_key)
            if cell is None:
                continue
            rows.append(
                [
                    employee.fullName,
                    departments,
                    expertises,
                    "Итого",
                    "",
                    "",
                    day_key,
                    _dec(cell.plannedHours),
                    _dec(cell.actualHours),
                    _dec(cell.capacityHours),
                    _dec(cell.availableHours),
                    cell.timeOffKind or "",
                    "Да" if cell.isWorkingDay else "Нет",
                ]
            )
            if not by_project:
                continue
            by_project_hours: dict[int, tuple[str, str, Decimal, Decimal]] = {}
            for item in cell.allocations:
                current = by_project_hours.get(item.projectId)
                if current is None:
                    by_project_hours[item.projectId] = (
                        item.requestNumber,
                        item.requestName,
                        item.plannedHours,
                        item.actualHours,
                    )
                else:
                    number, name, planned, actual = current
                    by_project_hours[item.projectId] = (
                        number,
                        name,
                        planned + item.plannedHours,
                        actual + item.actualHours,
                    )
            for project_id in sorted(by_project_hours, key=lambda pid: by_project_hours[pid][0]):
                number, name, planned, actual = by_project_hours[project_id]
                if planned == 0 and actual == 0:
                    continue
                rows.append(
                    [
                        employee.fullName,
                        departments,
                        expertises,
                        "Проект",
                        number,
                        name,
                        day_key,
                        _dec(planned),
                        _dec(actual),
                        _dec(cell.capacityHours),
                        _dec(cell.availableHours),
                        cell.timeOffKind or "",
                        "Да" if cell.isWorkingDay else "Нет",
                    ]
                )
    return rows


def export_workload_csv(
    workload: PlanningWorkloadOut,
    *,
    view_mode: str = "byProject",
    date_from: date,
    date_to: date,
) -> tuple[bytes, str]:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(_iter_export_rows(workload, view_mode=view_mode))
    # BOM for Excel on Windows
    content = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    stamp = datetime.now(MOSCOW_TZ).strftime("%Y%m%d")
    filename = f"planning-workload-{date_from.isoformat()}_{date_to.isoformat()}-{stamp}.csv"
    return content, filename


def export_workload_xlsx(
    workload: PlanningWorkloadOut,
    *,
    view_mode: str = "byProject",
    date_from: date,
    date_to: date,
) -> tuple[bytes, str]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Нагрузка"
    for col_index, title in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_index, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _TEXT_ALIGNMENT

    for row_index, values in enumerate(_iter_export_rows(workload, view_mode=view_mode), start=2):
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.alignment = _TEXT_ALIGNMENT

    last_row = max(1, sheet.max_row)
    for col_index in range(1, len(EXPORT_HEADERS) + 1):
        letter = get_column_letter(col_index)
        max_len = 12
        for row_index in range(1, last_row + 1):
            value = sheet.cell(row=row_index, column=col_index).value
            if value is None:
                continue
            max_len = max(max_len, min(48, len(str(value))))
        sheet.column_dimensions[letter].width = max_len + 2
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}{last_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    stamp = datetime.now(MOSCOW_TZ).strftime("%Y%m%d")
    filename = f"planning-workload-{date_from.isoformat()}_{date_to.isoformat()}-{stamp}.xlsx"
    return buffer.getvalue(), filename


__all__ = [
    "CSV_MEDIA_TYPE",
    "XLSX_MEDIA_TYPE",
    "export_workload_csv",
    "export_workload_xlsx",
]
